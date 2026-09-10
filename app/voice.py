import asyncio
import logging
import os
import time
from pathlib import Path

from app import crypto

log = logging.getLogger("voice")

VOICE_DIR = "voice_tmp"
Path(VOICE_DIR).mkdir(exist_ok=True)

# Bengali + English phrases that mean "send a voice message"
_VOICE_KW = [
    "voice dao", "voice de", "voice pathao", "voice ta dao", "voice ta pathao",
    "voice message", "voice note", "send voice", "send a voice", "voice reply",
    "audio message", "speak", "bolo", "voice a bolo", "voice chai", "voice lagbe",
]


def wants_voice(text: str) -> bool:
    if not text:
        return False
    t = text.lower()
    if any(kw in t for kw in _VOICE_KW):
        return True
    if "voice" in t and any(a in t for a in ["dao", "de", "pathao", "send", "chai", "lagbe", "note", "bolo"]):
        return True
    return False


def _clean(text: str) -> str:
    out = []
    for c in text:
        if ord(c) < 128 or c.isalpha() or c.isspace() or c in ".,!?-'":
            out.append(c)
    return " ".join("".join(out).split()).strip()


async def synth_via_provider(provider: dict, text: str) -> str | None:
    """Generic HTTP TTS: build the request from a stored config; return audio file path."""
    import json
    import base64
    import aiohttp
    try:
        voice = provider.get("voice") or ""
        endpoint = (provider.get("endpoint") or "").replace("{voice}", voice)
        tmpl = provider.get("body_template") or '{"text":"{text}"}'
        esc = json.dumps(text)[1:-1]  # JSON-escape into the template
        body_str = tmpl.replace("{text}", esc).replace("{voice}", voice)
        try:
            payload = json.loads(body_str)
        except Exception:
            payload = {"text": text}
        headers = {"Content-Type": "application/json"}
        hn, hv = provider.get("header_name"), provider.get("header_value_enc")
        if hn and hv:
            try:
                headers[hn] = crypto.decrypt(hv)
            except Exception:
                pass
        method = (provider.get("method") or "POST").upper()
        rt = (provider.get("response_type") or "audio").lower()
        path = os.path.join(VOICE_DIR, f"pv_{int(time.time()*1000)}.mp3")

        async with aiohttp.ClientSession(timeout=aiohttp.ClientTimeout(total=30)) as s:
            async with s.request(method, endpoint, json=payload, headers=headers) as r:
                if r.status != 200:
                    log.warning(f"voice provider HTTP {r.status}")
                    return None
                if rt == "audio":
                    with open(path, "wb") as f:
                        f.write(await r.read())
                elif rt in ("base64", "url"):
                    j = await r.json(content_type=None)
                    val = j
                    for key in (provider.get("json_path") or "").split("."):
                        if key:
                            val = val.get(key) if isinstance(val, dict) else None
                    if not val:
                        return None
                    if rt == "base64":
                        with open(path, "wb") as f:
                            f.write(base64.b64decode(val))
                    else:
                        async with s.get(val) as ar:
                            with open(path, "wb") as f:
                                f.write(await ar.read())
                else:
                    return None
        if os.path.exists(path) and os.path.getsize(path) > 0:
            return path
    except Exception as e:
        log.warning(f"voice provider error: {e}")
    return None


async def tts(text: str, voice: str, uid) -> str | None:
    try:
        import edge_tts
    except ImportError:
        return None
    script = _clean(text) or "Hello!"
    path = os.path.join(VOICE_DIR, f"v_{uid}_{int(time.time())}.mp3")
    try:
        await asyncio.wait_for(edge_tts.Communicate(script, voice).save(path), timeout=25.0)
        if os.path.exists(path) and os.path.getsize(path) > 0:
            return path
    except Exception as e:
        log.warning(f"tts error: {e}")
    return None


def extract_document_text(file_path: str) -> str:
    ext = Path(file_path).suffix.lower()
    try:
        if ext == ".pdf":
            import pdfplumber
            with pdfplumber.open(file_path) as pdf:
                return "\n".join((pg.extract_text() or "") for pg in pdf.pages[:6])[:3000]
        if ext in (".doc", ".docx"):
            from docx import Document
            return "\n".join(p.text for p in Document(file_path).paragraphs)[:3000]
        if ext in (".xlsx", ".xls"):
            import openpyxl
            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            rows = [str([c.value for c in row]) for row in wb.active.iter_rows(max_row=30)]
            return "\n".join(rows)[:2500]
        if ext == ".txt":
            return Path(file_path).read_text(encoding="utf-8", errors="ignore")[:3000]
    except Exception as e:
        log.warning(f"doc extract error: {e}")
    return ""
